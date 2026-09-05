"""Build `tests/fixtures/sample_cv.docx` from `tests/fixtures/cv_text.txt`.

The text file is Jim's real CV, one paragraph per line. This script recreates the
document shape described in docs/ARCHITECTURE.md ("Jim's CV specifics"):

- line 1 is the name, line 2 the `|`-separated contact line (the header);
- section headings are the all-caps lines PROFESSIONAL SUMMARY, PROFESSIONAL
  EXPERIENCE, CORE SKILLS & KNOWLEDGE, EDUCATION;
- bullets keep their literal `•    ` prefix (bullet + 4 spaces) as plain text, not
  Word list numbering;
- inside PROFESSIONAL EXPERIENCE, non-bullet lines are employer/date lines and
  project sub-headings;
- the skills section is a single `|`-delimited paragraph.

Every paragraph gets exactly one run with an explicit font so that formatting
preservation can be asserted in tests.

Run with: `uv run python tests/fixtures/build_fixture.py`
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

FIXTURES_DIR = Path(__file__).resolve().parent
CV_TEXT = FIXTURES_DIR / "cv_text.txt"
SAMPLE_CV = FIXTURES_DIR / "sample_cv.docx"

FONT_NAME = "Calibri"
BODY_SIZE = Pt(10)
NAME_SIZE = Pt(16)
HEADING_SIZE = Pt(11)
BULLET_PREFIX = "•    "

SECTION_HEADINGS = (
    "PROFESSIONAL SUMMARY",
    "PROFESSIONAL EXPERIENCE",
    "CORE SKILLS & KNOWLEDGE",
    "EDUCATION",
)


def _add(doc, text: str, *, size=BODY_SIZE, bold=False, italic=False, align=None,
         space_before=0, space_after=2):
    paragraph = doc.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    run.font.name = FONT_NAME
    run.font.size = size
    run.font.bold = bold
    run.font.italic = italic
    return paragraph


def build(src: Path = CV_TEXT, dst: Path = SAMPLE_CV) -> Path:
    lines = [line.rstrip("\r\n") for line in src.read_text(encoding="utf-8").splitlines()]
    lines = [line for line in lines if line.strip()]

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = BODY_SIZE
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Pt(36)
        section.left_margin = section.right_margin = Pt(43)

    current = "header"
    for index, line in enumerate(lines):
        if index == 0:
            _add(doc, line, size=NAME_SIZE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            continue
        if index == 1:
            _add(doc, line, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
            continue
        if line.upper() in SECTION_HEADINGS:
            current = line.upper()
            _add(doc, line, size=HEADING_SIZE, bold=True, space_before=6, space_after=3)
            continue
        if line.startswith("•"):
            _add(doc, line, space_after=1)
            continue
        if current == "PROFESSIONAL EXPERIENCE":
            # Employer/date line or project sub-heading.
            _add(doc, line, bold=True, space_before=3, space_after=1)
            continue
        _add(doc, line)

    doc.save(dst)
    return dst


if __name__ == "__main__":
    out = build()
    print(f"wrote {out}")
