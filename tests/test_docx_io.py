"""SCRUM-9: docx round-trip, table-cell paragraphs, formatting-preserving set_text."""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.shared import Pt
from lxml import etree

from resume_tailor.docx_io import Para, iter_paragraphs, load, save, set_text, structure_diff

HEADINGS = (
    "PROFESSIONAL SUMMARY",
    "PROFESSIONAL EXPERIENCE",
    "CORE SKILLS & KNOWLEDGE",
    "EDUCATION",
)
BULLET_PREFIX = "•    "


def _paras(path: Path) -> list[Para]:
    return iter_paragraphs(load(path))


def _first_bullet(paras: list[Para]) -> Para:
    return next(p for p in paras if p.literal_bullet_prefix == BULLET_PREFIX)


def _numbering_xml(para: Para) -> str:
    p_pr = para.paragraph._p.pPr
    if p_pr is None:
        return ""
    num_pr = p_pr.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr"
    )
    return "" if num_pr is None else etree.tostring(num_pr, encoding="unicode")


def test_fixture_has_expected_shape(sample_cv_path: Path) -> None:
    assert sample_cv_path.exists(), "run `uv run python tests/fixtures/build_fixture.py`"
    texts = [p.text for p in _paras(sample_cv_path)]
    for heading in HEADINGS:
        assert heading in texts, heading
    # Headings appear in CV order.
    positions = [texts.index(h) for h in HEADINGS]
    assert positions == sorted(positions)
    # Name and contact line come first.
    assert texts[0] == "MD RAWFUR MONZUR JIM"
    assert "|" in texts[1]


def test_iter_paragraphs_includes_table_cells() -> None:
    doc = Document()
    doc.add_paragraph("before table")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).paragraphs[0].text = "left cell"
    table.cell(0, 1).paragraphs[0].text = "right cell"
    doc.add_paragraph("after table")

    paras = iter_paragraphs(load(save(doc)))
    texts = [p.text for p in paras]
    assert texts == ["before table", "left cell", "right cell", "after table"]
    left = next(p for p in paras if p.text == "left cell")
    right = next(p for p in paras if p.text == "right cell")
    assert left.location == "body/table:0/row:0/col:0"
    assert right.location == "body/table:0/row:0/col:1"
    assert {p.id for p in paras} == {"p0", "p1", "p2", "p3"}


def test_word_and_char_budget(sample_cv_path: Path) -> None:
    paras = _paras(sample_cv_path)
    assert paras
    for para in paras:
        assert para.words == len(para.text.split()), para.id
        assert para.chars == len(para.text), para.id


def test_set_text_preserves_run_formatting(sample_cv_path: Path) -> None:
    doc = load(sample_cv_path)
    bullet = _first_bullet(iter_paragraphs(doc))
    run0 = bullet.paragraph.runs[0]
    before = (
        run0.font.name,
        run0.font.size,
        run0.font.bold,
        run0.font.italic,
        bullet.paragraph.style.name,
        _numbering_xml(bullet),
    )
    assert before[0] is not None and before[1] is not None  # fixture sets an explicit font

    # Give the paragraph a second, differently formatted run so we can see it removed.
    extra = bullet.paragraph.add_run(" EXTRA")
    extra.font.bold = True
    extra.font.size = Pt(20)
    assert len(bullet.paragraph.runs) == 2

    set_text(bullet, "Rewrote this bullet in the JD's vocabulary at the same length.")

    reloaded = load(save(doc))
    edited = next(p for p in iter_paragraphs(reloaded) if p.id == bullet.id)
    runs = edited.paragraph.runs
    assert len(runs) == 1
    after = (
        runs[0].font.name,
        runs[0].font.size,
        runs[0].font.bold,
        runs[0].font.italic,
        edited.paragraph.style.name,
        _numbering_xml(edited),
    )
    assert after == before
    assert edited.text == "Rewrote this bullet in the JD's vocabulary at the same length."
    assert bullet.words == len(bullet.text.split()) and bullet.chars == len(bullet.text)


def test_literal_bullet_prefix_preserved(sample_cv_path: Path) -> None:
    doc = load(sample_cv_path)
    bullet = _first_bullet(iter_paragraphs(doc))
    assert bullet.paragraph.text.startswith(BULLET_PREFIX)

    set_text(bullet, "New bullet body.")
    assert bullet.paragraph.text == BULLET_PREFIX + "New bullet body."
    assert bullet.text == "New bullet body."
    assert bullet.literal_bullet_prefix == BULLET_PREFIX

    # Passing a body that already carries a bullet must not double the prefix.
    set_text(bullet, BULLET_PREFIX + "Another body.")
    assert bullet.paragraph.text == BULLET_PREFIX + "Another body."

    reloaded = iter_paragraphs(load(save(doc)))
    assert next(p for p in reloaded if p.id == bullet.id).full_text == BULLET_PREFIX + "Another body."


def test_structure_diff_reports_only_text(sample_cv_path: Path) -> None:
    original = load(sample_cv_path)
    edited = load(sample_cv_path)
    paras = iter_paragraphs(edited)
    bullets = [p for p in paras if p.literal_bullet_prefix == BULLET_PREFIX]
    summary = next(p for p in paras if p.text.startswith("AI Engineer with years"))
    targets = [summary, bullets[3]]
    set_text(targets[0], "A rewritten professional summary.")
    set_text(targets[1], "A rewritten experience bullet.")

    diffs = structure_diff(original, load(save(edited)))
    assert [d.kind for d in diffs] == ["text", "text"]
    assert sorted(d.para_id for d in diffs) == sorted(t.id for t in targets)
    assert not [d for d in diffs if d.kind in {"style", "font", "numbering", "table"}]

    # Untouched documents have no diff at all.
    assert structure_diff(load(sample_cv_path), load(sample_cv_path)) == []


def test_roundtrip_untouched_is_identical(sample_cv_path: Path) -> None:
    doc = load(sample_cv_path)
    before = [(p.id, p.full_text, p.style, p.location, p.is_numbered) for p in iter_paragraphs(doc)]

    data = save(doc)
    assert isinstance(data, bytes)
    reloaded = load(data)
    after = [(p.id, p.full_text, p.style, p.location, p.is_numbered) for p in iter_paragraphs(reloaded)]

    assert after == before
    assert load(io.BytesIO(data)) is not None  # file-object input also works
    assert structure_diff(doc, reloaded) == []
